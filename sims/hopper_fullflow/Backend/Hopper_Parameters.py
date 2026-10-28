"""Parameter workbook: every input, component, assumption and open item in one
spreadsheet, generated from HOPPER_DATA.py so it always matches the model.

HOPPER_DATA.py stays the single source of truth: change a value there and
re-run this. The Status column is read from the comment next to each input
("GUESS" = nobody has confirmed it, "EDIT" = expected to change); Owner team
is a starting suggestion by section, for whoever keeps the numbers current
with the fluids, engines and injectors teams.

    python Backend/Hopper_Parameters.py      writes Hopper_Parameters.xlsx to PROP_RESULTS_DIRECTORY
"""

import ast
import datetime
import os
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import sys

if not __package__:   # run as a script (python Backend/Hopper_Parameters.py): make HOPPER_DATA and the Backend
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))   # package importable

import HOPPER_DATA as D
from Backend.Hopper_Hardware import COLUMNS as COMPONENT_COLUMNS, hardware_rows


OWNERS = {   # section header in HOPPER_DATA.py -> team that owns those numbers
    "RUN SETTINGS": "Sim", "THRUST CURVES": "GNC / trajectory", "FLUIDS": "Fluids",
    "PRESSURANT: COPVs": "Fluids", "TANKS": "Fluids", "PRESSURANT LEGS": "Fluids", "FEED LINES": "Fluids",
    "VENTS, RELIEFS, PURGES, DRAINS": "Fluids", "ENGINE": "Engines", "SEQUENCE": "Avionics / engines",
    "PLOTS": "Sim", "SOLVER NUMERICS": "Sim",
}

ASSUMPTIONS = [
    ("Liquids", "LOX and RP-1 are incompressible, with properties (ThermoProp) at the tank set pressure and "
                "loaded temperature. No propellant heating or boil-off during the burn."),
    ("Ullage collapse", "Collapse factor per tank: the COPV supplies the full pressurant flow, the ullage keeps "
                        "1 / collapse_factor of it. LOX 2.0, fuel 1.1: assumed values; Luca to supply "
                        "conservative and Vespula-heritage numbers."),
    ("Ullage gas", "Real-gas nitrogen (ThermoProp), solved in pressure-enthalpy. No heat transfer with the tank "
                   "walls; only the collapse factor removes gas."),
    ("COPV", "Real-gas nitrogen blowdown with no wall heat transfer (adiabatic), starting at 293 K."),
    ("Tanks", "Hydrostatic head and vehicle acceleration are neglected (as in vehicle_sim.py)."),
    ("Pressurant lines", "Each tube is a lumped equivalent CdA, A / sqrt(1 + f L / D) with the fully rough friction "
                         "factor, in series with the leg's valves and orifices (FullFlow SeriesCdA). No gas volume "
                         "or acoustic dynamics; adiabatic."),
    ("Gas restrictions", "FullFlow CompressibleOrifice (choked / subcritical) with the upstream real-gas R and gamma."),
    ("Liquid lines", "FullFlow DarcyWeisbach with flow inertia and Churchill friction; junctions have no volume."),
    ("Throttle valves", "OTV / FTV are ideal flow controllers: their CdA is solved to pass the thrust-target flow, "
                        "capped at their max CdA. Valve dynamics are a 0.01 s travel between profile steps."),
    ("Engine", "Equilibrium combustion map in (Pc, MR) times c* and Cf efficiencies that ramp linearly with Pc "
               "(c* 0.80 at 100 psia to 0.85 at 300 psia, Cf 0.94 at 100 to 0.97 at 250, flat outside; Luca's "
               "parameters, clamped instead of his max(1, ...)). In the transient each efficiency is taken at the "
               "previous step's Pc (one 0.05 s lag, under 0.1 % effect). Target MR 1.8. The chamber has no volume; propellant ignites as soon as it arrives."),
    ("Torch", "OIMV taps the LOX feed above OTV, FIMV the fuel feed above FTV (P&ID). Torch ox goes OIMV -> HX-1 "
              "-> torch injector, fuel FIMV -> torch injector, both into the main chamber, where they add to the "
              "chamber flow and MR. HX-1 is a pressure drop only (no heat exchange). All torch CdAs and the "
              "firing time (0.20-0.40 s) are GUESSes."),
    ("Start / shutdown", "GUESS sequence: OTV/FTV start opening at 0.30 s and take 0.20 s; the same ramp closes them "
                         "after the last profile point; 0.30 s of tail-off is simulated. The solver takes 0.01 s "
                         "steps through both ramps."),
    ("Bang-bang", "Each tank's solenoid closes above set + upper and opens below set - lower, sampled once per "
                  "accepted time step (0.05 s), so pressure overshoots the band slightly."),
    ("Regulator", "PR-REG is ideal: outlet = min(COPV, set) at the COPV gas temperature, no droop or lock-up."),
    ("Purges", "Purge GN2 is drawn from the COPV; OPV joins the LOX line below OTV, FPV the fuel line below FTV, "
               "TFPV the torch fuel line (P&ID). The gas only sees that back pressure; it is not added to the "
               "line or chamber flow. All purges are closed during the nominal burn."),
    ("Relief valves", "Opening rises linearly from crack to full-open dP (default 110 % of crack)."),
    ("Side paths", "Vent, relief, purge and drain flows only leave their source (no back-flow)."),
    ("Vent time", "Adiabatic blowdown of the ullage gas through the vent's CdA from set + upper limit to "
                  "1 psig, for a full tank (loaded ullage) and an empty tank. No LOX boil-off during the vent."),
    ("LOX drain", "OFD (same Jaksa valve as OVNT) plus any LOX_DRAIN_EXTRA_ELEMENTS as one CdA; ullage held at "
                  "each test pressure; outlet pressure = ullage + liquid head over a uniform cross-section "
                  "(tank volume / height). Incompressible LOX, no boil-off or flashing; the vented case is "
                  "optimistic because near-saturated LOX will partly flash in the valve."),
    ("Ambient", "Constant 14.8 psia (no altitude change during a hop)."),
    ("Hardware not modelled", "Quick disconnects (disconnected in flight), oxygen diffuser, torch spark "
                              "igniter, sensors."),
]

OPEN_ITEMS = [
    ("OBB / FBB CdA", "Model uses 18.68 mm2 for both. Master sheet: OBB = 2x PLV34 = 2.14 mm2, FBB = 1x PLV34 = "
                      "1.07 mm2. Which is right?"),
    ("Tank volumes", "Model 20 L LOX / 16 L fuel; master sheet 19.88 L (from Luca) / 15.18 L."),
    ("ORV sizing", "RegO PRV9432T550 has a 0.062 in orifice (about 1.95 mm2 geometric); the sheet requires "
                   "30.8 mm2 CdA (20.5 mm2 minimum). The chosen part looks undersized; check the RegO flow table."),
    ("FRV set pressure", "Sheet FRV tab says 525 psi; the part (PRV9432T550) is 550 psig. Model uses 550."),
    ("OVNT vent time", "With the Jaksa 9.4 mm2, an empty LOX tank takes about 64 s to vent from 505 psia to 1 psig "
                       "(requirement 60 s; sheet target CdA 14.1 mm2)."),
    ("Start-up tank spike", "Both tanks start at set pressure with the bang-bang solenoid open, so the first step "
                            "overshoots to about 517-520 psia. Start with bang_bang_initial_state = 0?"),
    ("Propellant residuals", "The 1dof curve uses about 10.8 kg LOX and 6.0 kg fuel of 19.43 / 9.7 kg loaded "
                             "(44 % / 38 % left). Are the loads or the curve out of date?"),
    ("Vendor flow data", "No published Cv found for Marotta PLV34, Swagelok SS-4R3A, Check-All BUDSSTF.125SS or the "
                         "UNOX 3/8 NO solenoid. Need datasheets or flow tests."),
    ("LOX drain path", "Drain time uses only the OFD valve (9.4 mm2). OQD and the GSE drain line will slow it; add "
                       "them to LOX_DRAIN_EXTRA_ELEMENTS. Drain exit height below the tank and the drain ullage "
                       "pressure are GUESSes."),
    ("Nozzle note", "LinearizedNozzle docstring says Pc = Pa + 1 psi; NOZZLE_LINEAR_PRESSURE_DROP is 0.5 psi."),
    ("Collapse factor", "LOX 2.0 / fuel 1.1 are assumed. Waiting on Luca for conservative and Vespula values."),
]


def _sections_and_comments():
    """{line number: section}, {line number: comment} for HOPPER_DATA.py."""
    lines = Path(D.__file__).read_text(encoding="utf-8").splitlines()
    sections, comments, section = {}, {}, ""
    for number, line in enumerate(lines, start=1):
        match = re.match(r"#\s*----\s*(.+?)\s*----\s*#", line.strip())
        if match:
            section = match.group(1)
        sections[number] = section
        if "#" in line and not line.lstrip().startswith("#"):
            comments[number] = line.split("#", 1)[1].strip()
    return lines, sections, comments


def _status(text):
    upper = text.upper()
    return "GUESS" if "GUESS" in upper else "EDIT" if "EDIT" in upper else ""


def _owner(section):
    for key, owner in OWNERS.items():
        if section.upper().startswith(key):
            return owner
    return ""


def _format(value):
    if isinstance(value, float):
        return float(f"{value:.6g}")
    if isinstance(value, (int, str, bool)):
        return value
    return str(value)


def input_rows():
    """[section, name, as written, value (SI), note, status, owner] for every
    plain input in HOPPER_DATA.py, with the Tank objects expanded."""
    lines, sections, comments = _sections_and_comments()
    tree = ast.parse("\n".join(lines))
    rows = []
    for node in tree.body:
        if not isinstance(node, ast.Assign) or not isinstance(node.targets[0], ast.Name):
            continue
        name = node.targets[0].id
        if name.startswith("_") or not name.isupper():
            continue
        value = getattr(D, name)
        section = sections.get(node.lineno, "")
        note = " ".join(comments.get(n, "") for n in range(node.lineno, node.end_lineno + 1)).strip()
        if isinstance(value, list) or name in ("NOT_MODELLED", "BRANCHES", "REGULATORS"):
            continue   # element lists are on the Components sheet
        if isinstance(node.value, ast.Call) and getattr(node.value.func, "id", "") == "Tank":
            for keyword in node.value.keywords:
                line_note = comments.get(keyword.value.lineno, "")
                rows.append([section, f"{name}.{keyword.arg}", ast.get_source_segment("\n".join(lines), keyword.value),
                             _format(getattr(value, keyword.arg)), line_note, _status(line_note), _owner(section)])
            continue
        if isinstance(value, dict) and name != "VENT_TIME_REQUIREMENTS":
            written = ast.get_source_segment("\n".join(lines), node.value)
            rows.append([section, name, written if len(written) < 200 else "(see HOPPER_DATA.py)",
                         f"{len(value)} entries", note, _status(note), _owner(section)])
            continue
        written = ast.get_source_segment("\n".join(lines), node.value)
        rows.append([section, name, written, _format(value), note, _status(note), _owner(section)])
    return rows


def _sheet(workbook, title, header, rows, widths):
    sheet = workbook.create_sheet(title)
    sheet.append(header)
    for row in rows:
        sheet.append(list(row))
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    for cell in sheet[1]:
        cell.font, cell.fill = bold, fill
    guess = PatternFill("solid", fgColor="FFF2CC")
    status_column = header.index("Status") if "Status" in header else None
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if status_column is not None and row[status_column].value == "GUESS":
            for cell in row:
                cell.fill = guess
    for k, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(k)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def write_workbook(directory=D.PROP_RESULTS_DIRECTORY):
    os.makedirs(directory, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    today = datetime.date.today().isoformat()

    inputs = [row + ["", ""] for row in input_rows()]
    _sheet(workbook, "Inputs",
           ["Section", "Name", "As written (units)", "Value (SI)", "Note / source", "Status", "Owner team",
            "Owner", "Last confirmed"], inputs, [22, 34, 34, 14, 60, 9, 16, 14, 14])

    components = []
    for row in hardware_rows():
        status = _status(row[-1])
        path = row[0].lower()
        team = "Injectors" if row[2] == "Injector" else "Engines" if "regen" in row[1].lower() else "Fluids"
        components.append(row + [status, team if path != "not in the flow model" else "", "", ""])
    _sheet(workbook, "Components", COMPONENT_COLUMNS + ["Status", "Owner team", "Owner", "Last confirmed"],
           components, [40, 22, 18, 10, 46, 52, 9, 12, 14, 14])

    _sheet(workbook, "Assumptions", ["Topic", "Assumption"], ASSUMPTIONS, [22, 120])
    guesses = [(row[1], f"{row[4]} (value {row[3]})") for row in inputs if row[5] == "GUESS"]
    guesses += [(row[1], f"{row[5]} (CdA {row[3]} mm2)") for row in hardware_rows() if _status(row[-1]) == "GUESS"]
    _sheet(workbook, "Open items", ["Item", "Detail"],
           OPEN_ITEMS + [(f"GUESS: {name}", detail) for name, detail in guesses], [30, 120])

    about = workbook.create_sheet("About", 0)
    for line in [f"Hopper FullFlow parameters, generated {today} from HOPPER_DATA.py by Hopper_Parameters.py.",
                 "HOPPER_DATA.py is the source of truth: edit values there, then re-run this script.",
                 "Status GUESS (yellow rows) = not confirmed by the owning team. EDIT = expected to change.",
                 "Sheets: Inputs, Components (every P&ID item), Assumptions, Open items."]:
        about.append([line])
    about.column_dimensions["A"].width = 110

    path = os.path.join(directory, "Hopper_Parameters.xlsx")
    workbook.save(path)
    return path


if __name__ == "__main__":
    print("wrote", write_workbook())
